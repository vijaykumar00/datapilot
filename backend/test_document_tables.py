"""Tables from PDF and Word uploads become datasets (one 'sheet' per table)."""

import io
import zipfile
from pathlib import Path

import pandas as pd
import pytest

from core.document_tables import DocumentError, extract_document_tables
from core.file_manager import FileManager
from core.parsing import ParseError, parse_file
from core.question_engine import plan_question
from core.data_store import execute_select

reportlab = pytest.importorskip("reportlab")
pytest.importorskip("pdfplumber")

from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import getSampleStyleSheet  # noqa: E402
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle  # noqa: E402

GRID = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])


def _pdf(path: Path, story) -> Path:
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
    return path


def _sales_rows(n):
    regions = ["North", "South", "East", "West"]
    return [[f"2026-{(i % 12) + 1:02d}-01", regions[i % 4], f"${1000 + i * 10:,.2f}", str(i % 5 + 1)] for i in range(n)]


@pytest.fixture
def statement_pdf(tmp_path):
    """A 'bank statement': title text, one table running over three pages, then a second table."""
    styles = getSampleStyleSheet()
    header = ["Date", "Region", "Revenue", "Units"]
    rows = _sales_rows(120)
    story = [Paragraph("Quarterly statement", styles["Title"]),
             Table([header] + rows, repeatRows=1, style=GRID),  # repeats its header on every page
             PageBreak(),
             Paragraph("Fees", styles["Heading2"]),
             Table([["Fee", "Amount"], ["Service", "12.50"], ["Wire", "30.00"]], style=GRID)]
    return _pdf(tmp_path / "statement.pdf", story), rows


def test_pdf_table_over_several_pages_becomes_one_dataset(statement_pdf):
    path, rows = statement_pdf
    df, meta = parse_file(path, ".pdf")
    assert meta["document_type"] == "pdf"
    assert len(meta["sheet_names"]) == 2
    first, second = meta["sheet_names"]
    assert first.startswith("Table 1 (pages 1-") and second.startswith("Table 2 (page ")
    assert list(df.columns) == ["Date", "Region", "Revenue", "Units"]
    assert len(df) == len(rows)  # repeated header rows were dropped, no rows lost
    # "$1,000.00" text becomes a number, so totals work.
    assert df["Revenue"].dtype.kind == "f"
    assert df["Revenue"].sum() == pytest.approx(sum(1000 + i * 10 for i in range(len(rows))))
    assert meta["sheet_columns"][second] == ["Fee", "Amount"]


def test_switching_to_a_table_reads_only_its_pages(statement_pdf):
    path, _ = statement_pdf
    _, meta = parse_file(path, ".pdf")
    second = meta["sheet_names"][1]
    locator = meta["table_locators"][second]
    df, _ = parse_file(path, ".pdf", locator)
    assert list(df.columns) == ["Fee", "Amount"] and df["Amount"].sum() == pytest.approx(42.5)
    # The first table re-read by locator is the same as from the full scan.
    full, _ = parse_file(path, ".pdf")
    again, _ = parse_file(path, ".pdf", meta["table_locators"][meta["sheet_names"][0]])
    pd.testing.assert_frame_equal(full, again)
    # Selecting by display name also works (slower path: whole document).
    by_name, _ = parse_file(path, ".pdf", second)
    pd.testing.assert_frame_equal(by_name, df)


def test_question_engine_answers_on_pdf_table(statement_pdf):
    path, rows = statement_pdf
    df, _ = parse_file(path, ".pdf")
    plan = plan_question("total revenue by region", df, "t")
    result = execute_select(plan.sql, {"t": df}).rows
    expected = pd.DataFrame(rows, columns=["d", "region", "rev", "u"])
    expected["rev"] = expected["rev"].str.replace("[$,]", "", regex=True).astype(float)
    assert {r["region"]: r["total_revenue"] for r in result} == pytest.approx(
        expected.groupby("region").rev.sum().to_dict())


def test_text_only_and_scanned_pdfs_explain_why(tmp_path):
    styles = getSampleStyleSheet()
    text_pdf = _pdf(tmp_path / "letter.pdf", [Paragraph("Dear customer, " * 40, styles["Normal"])])
    with pytest.raises(ParseError, match="No tables were found"):
        parse_file(text_pdf, ".pdf")

    from reportlab.pdfgen import canvas
    scanned = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(scanned))
    c.rect(50, 50, 400, 600, fill=1)  # a page with only graphics, like a scanned image
    c.save()
    with pytest.raises(ParseError, match="looks scanned"):
        parse_file(scanned, ".pdf")


def test_page_limit(tmp_path, monkeypatch):
    path = _pdf(tmp_path / "long.pdf", [Table([["a", "b"], ["1", "2"]], style=GRID), PageBreak(),
                                         Table([["a", "b"], ["3", "4"]], style=GRID)])
    monkeypatch.setenv("MAX_DOCUMENT_PAGES", "1")
    with pytest.raises(DocumentError, match="has 2 pages"):
        extract_document_tables(path, ".pdf")


def test_word_tables(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Price list")
    t = doc.add_table(rows=1, cols=3)
    for cell, text in zip(t.rows[0].cells, ["Product", "Price", "Stock"]):
        cell.text = text
    for name, price, stock in [("Widget", "9.99", "120"), ("Gadget", "1,299.00", "8")]:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, price, stock
    doc.add_table(rows=1, cols=1).rows[0].cells[0].text = "Signature"  # layout box, not data
    path = tmp_path / "prices.docx"
    doc.save(path)

    df, meta = parse_file(path, ".docx")
    assert meta["sheet_names"] == ["Table 1"] and meta["document_type"] == "docx"
    assert df.to_dict("records") == [{"Product": "Widget", "Price": 9.99, "Stock": 120},
                                     {"Product": "Gadget", "Price": 1299.0, "Stock": 8}]

    empty = docx.Document()
    empty.add_paragraph("Only text")
    empty.save(tmp_path / "text.docx")
    with pytest.raises(ParseError, match="No tables were found in this Word file"):
        parse_file(tmp_path / "text.docx", ".docx")


def test_upload_validation_for_documents():
    fm = FileManager()
    assert fm._validate_extension("report.PDF") == ".pdf"
    assert fm._validate_extension("prices.docx") == ".docx"
    with pytest.raises(ValueError):
        fm._validate_extension("old.doc")
    fm._validate_magic(b"%PDF-1.7\n...", ".pdf")
    with pytest.raises(ValueError, match="PDF"):
        fm._validate_magic(b"<html>not a pdf", ".pdf")

    def zipped(name):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr(name, "<x/>")
        return buf.getvalue()

    fm._validate_magic(zipped("word/document.xml"), ".docx")
    fm._validate_workbook_expansion(zipped("word/document.xml"), ".docx")
    # A spreadsheet (or any other zip) renamed to .docx is rejected, and vice versa.
    with pytest.raises(ValueError, match="Word document"):
        fm._validate_workbook_expansion(zipped("xl/workbook.xml"), ".docx")
    with pytest.raises(ValueError, match="Excel workbook"):
        fm._validate_workbook_expansion(zipped("word/document.xml"), ".xlsx")


def test_pdf_upload_end_to_end_and_table_switch(statement_pdf):
    """Upload → background ingest (killable subprocess) → switch table → workspace isolation."""
    import os
    import uuid

    path, rows = statement_pdf
    ws = f"ws_{uuid.uuid4().hex[:8]}"
    fm = FileManager()
    dataset_id = fm.stage_upload(path, "statement.pdf", os.path.getsize(path), workspace_id=ws, user_id="u1")
    try:
        summary = fm.ingest(dataset_id)
        assert summary["row_count"] == len(rows)
        meta = summary["metadata"]
        assert meta["document_type"] == "pdf" and len(meta["sheet_names"]) == 2 and meta["table_locators"]

        switched = fm.switch_sheet(dataset_id, meta["sheet_names"][1], ws)
        assert switched["row_count"] == 2 and switched["metadata"]["active_sheet"] == meta["sheet_names"][1]
        record = FileManager().get_record(dataset_id, ws)  # fresh process view
        assert list(record.df.columns) == ["Fee", "Amount"]
        assert FileManager().get_record(dataset_id, "another_workspace") is None
    finally:
        fm.delete_file(dataset_id)


def test_unreadable_pdf_upload_fails_with_clear_message(tmp_path):
    import os
    import uuid

    styles = getSampleStyleSheet()
    path = _pdf(tmp_path / "memo.pdf", [Paragraph("Just a memo. " * 30, styles["Normal"])])
    ws = f"ws_{uuid.uuid4().hex[:8]}"
    fm = FileManager()
    dataset_id = fm.stage_upload(path, "memo.pdf", os.path.getsize(path), workspace_id=ws, user_id="u1")
    try:
        with pytest.raises(ValueError, match="No tables were found in this PDF"):
            fm.ingest(dataset_id)
    finally:
        fm.delete_file(dataset_id)
