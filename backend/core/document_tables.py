"""
document_tables.py — pull the tables out of PDF and Word files.

Each table found becomes one "sheet" of the upload, so the rest of Tallis
(sheet switching, the question engine, reports) works on it unchanged.

* PDF: pdfplumber (MIT) finds ruled and well-aligned tables page by page.  A
  table that runs over several pages (with or without a repeated header row) is
  joined back into one table.  Pages are capped (MAX_DOCUMENT_PAGES).
* Word (.docx): python-docx tables, in document order.

Text-only and scanned documents raise ``DocumentError`` with a message that can
be shown to the user as-is.  This runs inside the killable parse subprocess, so
a pathological file is bounded by FILE_PARSE_TIMEOUT_SECONDS.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

DOCUMENT_EXTENSIONS = {".pdf", ".docx"}
MAX_TABLES = 100  # tables offered per document (the first ones, in reading order)
_LOCATOR_RE = re.compile(r"^@p(\d+)-(\d+)#(\d+)$")
_NUMERIC_CELL = re.compile(r"^[\s$€£¥₹(+-]*\d[\d\s.,]*%?\)?\s*$")


class DocumentError(ValueError):
    """A user-facing reason why no data could be taken from a document."""


def max_document_pages() -> int:
    try:
        return max(1, int(os.getenv("MAX_DOCUMENT_PAGES", "300")))
    except ValueError:
        return 300


def count_pdf_pages(path: str | Path) -> int | None:
    """Page count without parsing page content (fast; None if the file can't be read)."""
    try:
        import pypdfium2

        doc = pypdfium2.PdfDocument(str(path))
        try:
            return len(doc)
        finally:
            doc.close()
    except Exception:
        return None


@dataclass
class _Table:
    start_page: int  # 1-based
    index: int  # position among the raw tables on its start page
    rows: list[list[str]]
    end_page: int = 0
    header: list[str] = field(default_factory=list)

    @property
    def locator(self) -> str:
        return f"@p{self.start_page}-{self.end_page}#{self.index}"

    def name(self, number: int) -> str:
        pages = f"page {self.start_page}" if self.start_page == self.end_page else f"pages {self.start_page}-{self.end_page}"
        return f"Table {number} ({pages})" if self.start_page else f"Table {number}"


def _clean_cell(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _clean_rows(rows) -> list[list[str]]:
    out = [[_clean_cell(c) for c in row] for row in rows if row]
    out = [row for row in out if any(row)]
    if not out:
        return []
    width = max(len(r) for r in out)
    out = [r + [""] * (width - len(r)) for r in out]
    # Drop columns that are empty in every row (layout artefacts of merged cells).
    keep = [i for i in range(width) if any(r[i] for r in out)]
    return [[r[i] for i in keep] for r in out]


def _looks_like_data_row(row: list[str]) -> bool:
    filled = [c for c in row if c]
    return bool(filled) and sum(bool(_NUMERIC_CELL.match(c)) for c in filled) / len(filled) >= 0.5


def _is_real_table(rows: list[list[str]]) -> bool:
    # A header plus at least one row, two or more columns: anything smaller is
    # usually a boxed paragraph or a form field, not data.
    return len(rows) >= 2 and len(rows[0]) >= 2


def _merge_continuations(tables: list[_Table]) -> list[_Table]:
    """Join a table with the tables that continue it on the following pages."""
    merged: list[_Table] = []
    for table in tables:
        table.end_page = table.start_page
        table.header = table.rows[0]
        prev = merged[-1] if merged else None
        continues = (
            prev is not None
            and table.start_page == prev.end_page + 1
            and table.index == 0
            and len(table.rows[0]) == len(prev.header)
            and (table.rows[0] == prev.header or _looks_like_data_row(table.rows[0]))
        )
        if continues:
            body = table.rows[1:] if table.rows[0] == prev.header else table.rows
            prev.rows.extend(body)
            prev.end_page = table.start_page
        else:
            merged.append(table)
    return merged


def _raw_pdf_tables(path: Path, first: int, last: int | None) -> tuple[list[_Table], int, bool]:
    """Return (raw tables, page count, any text found) for pages first..last (1-based)."""
    import pdfplumber

    tables: list[_Table] = []
    has_text = False
    try:
        pdf = pdfplumber.open(str(path))
    except Exception as exc:
        raise DocumentError("This PDF could not be opened. It may be damaged or password-protected.") from exc
    with pdf:
        total = len(pdf.pages)
        limit = max_document_pages()
        if last is None and total > limit:
            raise DocumentError(
                f"This PDF has {total} pages. Tallis reads tables from documents of up to {limit} pages; "
                "split the file or export just the pages with tables."
            )
        stop = min(total, last if last is not None else total)
        for page_no in range(first, stop + 1):
            page = pdf.pages[page_no - 1]
            try:
                if not has_text and page.chars:
                    has_text = True
                found = page.extract_tables()
            except Exception:
                continue
            finally:
                page.flush_cache()
            kept = 0
            for raw in found:
                rows = _clean_rows(raw)
                if _is_real_table(rows):
                    tables.append(_Table(start_page=page_no, index=kept, rows=rows))
                    kept += 1
    return tables, total, has_text


def _docx_tables(path: Path) -> list[_Table]:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise DocumentError("This Word file could not be opened. Re-save it as .docx and try again.") from exc
    tables: list[_Table] = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = []
            last = None
            for cell in row.cells:
                # Horizontally merged cells repeat the same cell object; keep one copy.
                if cell._tc is last:
                    continue
                last = cell._tc
                cells.append(cell.text)
            rows.append(cells)
        rows = _clean_rows(rows)
        if _is_real_table(rows):
            tables.append(_Table(start_page=0, index=len(tables), rows=rows, end_page=0, header=rows[0]))
    return tables


def _to_frame(table: _Table) -> pd.DataFrame:
    return pd.DataFrame(table.rows[1:], columns=table.rows[0]).replace("", None)


def extract_document_tables(path: str | Path, ext: str, sheet: str | None = None) -> tuple[pd.DataFrame, dict]:
    """Return (frame of the selected table, parse metadata) for a PDF or Word file.

    *sheet* is either a table name returned earlier in ``sheet_names`` or the
    locator stored for it (``@p<start>-<end>#<index>``).  With a locator only the
    pages that table spans are read again, so switching tables in a long PDF is
    fast.
    """
    path = Path(path)
    ext = ext.lower()
    locator = _LOCATOR_RE.match(sheet or "")
    if ext == ".pdf" and locator:
        start, end, index = (int(g) for g in locator.groups())
        raw, _, _ = _raw_pdf_tables(path, start, end)
        for table in _merge_continuations([t for t in raw if t.start_page > start or t.index >= index]):
            if table.start_page == start and table.index == index:
                return _to_frame(table), {"document_type": "pdf"}
        raise DocumentError("That table could not be found in the document any more.")

    if ext == ".pdf":
        raw, pages, has_text = _raw_pdf_tables(path, 1, None)
        tables = _merge_continuations(raw)
        if not tables:
            if not has_text:
                raise DocumentError(
                    "This PDF looks scanned (it contains images of pages, not text), so its tables can't be read yet. "
                    "Upload the original Excel/CSV, or a PDF exported from the source program."
                )
            raise DocumentError(
                "No tables were found in this PDF. Tallis reads tables (rows and columns) from documents; "
                "plain text pages aren't supported yet."
            )
    elif ext == ".docx":
        tables, pages = _docx_tables(path), None
        if not tables:
            raise DocumentError(
                "No tables were found in this Word file. Tallis reads tables (rows and columns) from documents; "
                "plain text isn't supported yet."
            )
    else:
        raise DocumentError(f"Unsupported document type '{ext}'.")

    truncated = len(tables) > MAX_TABLES
    tables = tables[:MAX_TABLES]
    names = [t.name(i + 1) for i, t in enumerate(tables)]
    chosen = names.index(sheet) if sheet in names else 0
    meta = {
        "document_type": ext.lstrip("."),
        "sheet_names": names,
        "active_sheet": names[chosen],
        "header_row": 0,
        "sheet_columns": {n: [c for c in t.header if c] for n, t in zip(names[:25], tables[:25])},
        # How to re-read a table later without scanning the whole document.
        "table_locators": {n: t.locator for n, t in zip(names, tables)} if ext == ".pdf" else {},
    }
    if truncated:
        meta["tables_truncated"] = MAX_TABLES
    if pages is not None:
        meta["page_count"] = pages
    return _to_frame(tables[chosen]), meta
